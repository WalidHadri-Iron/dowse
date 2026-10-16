"""Builds a folder of test files for Dowse and its eval queries (evals/testbed.json; its hard/ part, evals/hard.json).

    uv run python scripts/make_testbed.py [folder]     (default: %LOCALAPPDATA%\\Dowse\\Testbed)

Photos get camera names (IMG_1001.jpg), so a search finds them by what they show, never by their names. Traps: a code
project (only its screenshot counts), a private key, and a scanned PDF only OCR can read.
"""

import io
import subprocess
import sys
import tempfile
import time
import wave
from fractions import Fraction
from pathlib import Path

import av
import httpx
import numpy as np
from docx import Document
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from dowse import paths  # noqa: E402
import hard_testbed  # noqa: E402  (beside this script)

# (camera name, Wikimedia Commons file)
PHOTOS = [
    ("IMG_1001.jpg", "Plains_Zebra_Equus_quagga.jpg"),
    ("IMG_1002.jpg", "Golden_Retriever_Dukedestiny01_drvd.jpg"),
    ("IMG_1003.jpg", "Tour_Eiffel_Wikimedia_Commons.jpg"),
    ("IMG_1004.jpg", "Pizza-3007395.jpg"),
    ("IMG_1005.jpg", "Laptop-hard-drive-exposed.jpg"),
    ("IMG_1006.jpg", "Emperor_Penguin_Manchot_empereur.jpg"),
    ("IMG_1007.jpg", "Sunflower_sky_backdrop.jpg"),
    ("IMG_1008.jpg", "Lightning_over_Oradea_Romania_zoom.jpg"),
    ("IMG_1009.jpg", "Violin_VL100.png"),
    ("DSC_2001.jpg", "Red_Apple.jpg"),
    ("DSC_2002.jpg", "Snow-capped_mountain_peak.jpg"),
    ("DSC_2003.jpg", "Cat_November_2010-1a.jpg"),
    ("DSC_2004.jpg", "Bicycle_diagram-unif.svg"),
    ("PXL_3001.jpg", "Hamburger_(black_bg).jpg"),
    ("PXL_3002.jpg", "Crested_Saguaro_cactus.jpg"),
]

SCREENSHOTS = {
    "Screenshot 2026-09-14 101500.png": ["Checkout", "Payment declined", "Your card was declined. Error code E-4021.",
                                        "Try another payment method or contact your bank."],
    "Screenshot 2026-09-20 183012.png": ["PowerShell", "PS C:\\dev\\app> python train.py",
                                        "Traceback (most recent call last):",
                                        "ModuleNotFoundError: No module named 'torch'"],
    "Screenshot (12).png": ["Booking confirmed", "Flight LH 716 Frankfurt -> Tokyo Haneda",
                            "Departure 3 November 2026, 13:55", "Booking reference K7Q2LM"],
}

DOCS = {
    "notes/banana bread.txt": "Banana bread\n\nMash three ripe bananas, mix with melted butter, sugar, one egg and "
                              "flour. Bake for 60 minutes at 175 C until a skewer comes out clean.",
    "notes/kyoto trip.md": "# Kyoto in November\n\nDay 1: Fushimi Inari shrine early in the morning, before the crowds."
                           " Day 2: the bamboo grove in Arashiyama, then a boat on the river. Day 3: tea ceremony "
                           "in Gion and dinner at a small izakaya.",
    "notes/electric cars.html": "<html><head><title>Range anxiety</title></head><body><h1>Range anxiety</h1><p>Most "
                                "electric cars now drive more than 400 km on one charge, and fast chargers along "
                                "motorways refill 80% of the battery in about 25 minutes.</p></body></html>",
    "notes/note-fr.txt": "Rendez-vous chez le dentiste mardi à 9 heures pour un détartrage. Apporter la carte vitale.",
    "notes/bn-note.txt": "আগামী শুক্রবার বাজার থেকে ইলিশ মাছ কিনতে হবে। মা সর্ষে ইলিশ রান্না করবেন।",
    "mail/dentist.eml": "From: clinic@example.org\nTo: me@example.org\nSubject: Appointment reminder\n"
                        "Content-Type: text/plain; charset=utf-8\n\nThis is a reminder of your dental check-up on "
                        "Tuesday at 9:00. Please arrive ten minutes early.\n",
    "code/webapp/package.json": '{"name": "webapp"}',
    "code/webapp/README.md": "# Webapp\n\nThe tomato sauce recipe is NOT here; this file must never be indexed.",
    ".ssh/id_rsa": "-----BEGIN OPENSSH PRIVATE KEY-----\nnot a real key\n-----END OPENSSH PRIVATE KEY-----\n",
    "secrets/id_ed25519": "not a real key",
}

LEASE = [
    ("Residential lease agreement", "This lease is made between the landlord and the tenant for the apartment at "
     "12 Garden Street. The term is twelve months starting 1 January 2027. Rent is 1,200 euros per month, due on the "
     "first day of each month, paid by bank transfer. " * 2),
    ("Section 7: Pets", "The tenant may keep one cat or one small dog in the apartment with the landlord's written "
     "consent. Any damage caused by pets will be deducted from the security deposit. Exotic animals, reptiles and "
     "birds of prey are not allowed. " * 2),
    ("Section 9: Termination", "Either party may end this agreement with three months' written notice. The deposit "
     "is returned within thirty days after the keys are handed back and the apartment is inspected. " * 2),
]

INVOICE = ("Invoice INV-2026-0042", "Acme Garden Supplies. Billed to: Alex. Items: 2 terracotta pots, 1 bag of "
           "potting soil, 1 watering can. Total due: 86.40 euros, payable within 14 days. Thank you for your order. "
           "Our bank details are printed on the back of this invoice. " * 2)

SCAN_TEXT = ["MEDICAL CERTIFICATE", "", "The patient was examined on 2 October 2026", "and is fit to return to work",
             "from 6 October 2026.", "", "Dr. A. Rahman, General Practitioner"]


# The podcast's topics, in order; their moments go into evals/testbed.json.
PODCAST = [
    ("Welcome back to the show. This week in technology news: a new phone with a folding screen, faster chips for "
     "laptops, and a big update for a popular messaging app that adds end to end encryption for group calls."),
    ("Now let's talk about sleep. Most adults need seven to nine hours of sleep every night. Going to bed at the "
     "same time each day keeps your body clock steady, and screens before bedtime make it harder to fall asleep."),
    ("In the kitchen this week: a simple tomato sauce for pasta. Fry garlic in olive oil, add crushed tomatoes and "
     "fresh basil, season with salt, and let it simmer for twenty minutes before tossing it with spaghetti."),
    ("And finally, sports. Last night's football match ended two to one after a late goal in the ninetieth minute. "
     "The coach praised the goalkeeper, who saved a penalty in the first half."),
]
MEMO = "Remember to call the plumber about the leaking kitchen sink on Thursday morning, before nine."
SOUNDTRACK = ("Welcome to the wildlife channel. Today we visit the savanna, then a city in Europe, and we end the day "
              "with dinner at a small Italian restaurant.")


def speak(text: str, scratch: Path) -> tuple[np.ndarray, int]:
    """Windows' own speech synthesizer (System.Speech): mono int16 samples and their rate."""
    out = scratch / f"tts{abs(hash(text))}.wav"
    script = ("Add-Type -AssemblyName System.Speech; $s = New-Object System.Speech.Synthesis.SpeechSynthesizer; "
              f"$s.SetOutputToWaveFile('{out}'); $s.Speak([Console]::In.ReadToEnd()); $s.Dispose()")
    subprocess.run(["powershell", "-NoProfile", "-Command", script], input=text, text=True, check=True,
                   creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    with wave.open(str(out)) as clip:
        samples = np.frombuffer(clip.readframes(clip.getnframes()), np.int16)
        rate = clip.getframerate()
    out.unlink()
    return samples, rate


def encode_audio(samples: np.ndarray, rate: int, target: Path, codec: str) -> None:
    with av.open(str(target), "w") as container:
        stream = container.add_stream(codec, rate=rate, layout="mono")
        _mux_audio(container, stream, samples, rate)


def _mux_audio(container, stream, samples: np.ndarray, rate: int) -> None:
    frame_size = stream.codec_context.frame_size or 1024
    resampler = av.AudioResampler(format=stream.codec_context.format.name, layout="mono", rate=rate)
    for start in range(0, len(samples), frame_size):
        frame = av.AudioFrame.from_ndarray(samples[None, start:start + frame_size], format="s16", layout="mono")
        frame.sample_rate, frame.pts, frame.time_base = rate, start, Fraction(1, rate)
        for converted in resampler.resample(frame):
            container.mux(stream.encode(converted))
    container.mux(stream.encode(None))


def make_podcast(target: Path, scratch: Path) -> list[tuple[float, float]]:
    """The topics one after another with two seconds of quiet between them; returns each topic's (start, end)."""
    pieces, moments, rate, at = [], [], 22050, 0.0
    for text in PODCAST:
        samples, rate = speak(text, scratch)
        moments.append((round(at, 1), round(at + len(samples) / rate, 1)))
        pieces += [samples, np.zeros(2 * rate, np.int16)]
        at += len(samples) / rate + 2
    encode_audio(np.concatenate(pieces), rate, target, "libmp3lame")
    return moments


def make_slideshow(photos: list[Path], target: Path, scratch: Path, seconds: int = 6, fps: int = 10) -> None:
    """Each photo for `seconds`, with a spoken soundtrack. x264's default keyframe interval (250 frames) leaves one
    keyframe in the clip: the reader has to seek into the gap to see the second and third photo."""
    voice, rate = speak(SOUNDTRACK, scratch)
    with av.open(str(target), "w") as container:
        video = container.add_stream("libx264", rate=fps)
        video.width, video.height, video.pix_fmt = 960, 540, "yuv420p"
        audio = container.add_stream("aac", rate=rate, layout="mono")
        index = 0
        for photo in photos:
            image = Image.open(photo).convert("RGB")
            image.thumbnail((960, 540))
            canvas = Image.new("RGB", (960, 540))
            canvas.paste(image, ((960 - image.width) // 2, (540 - image.height) // 2))
            for _ in range(seconds * fps):
                frame = av.VideoFrame.from_image(canvas)
                frame.pts = index
                container.mux(video.encode(frame))
                index += 1
        container.mux(video.encode(None))
        _mux_audio(container, audio, voice, rate)


# Public-domain books (Project Gutenberg): long files, read a step at a time, that every query is scored against.
BOOKS = {"books/pride and prejudice.txt": 1342, "books/sherlock holmes.pdf": 1661}


def gutenberg(number: int) -> str:
    text = httpx.get(f"https://www.gutenberg.org/cache/epub/{number}/pg{number}.txt", follow_redirects=True,
                     timeout=60).text
    start, end = text.find("*** START OF"), text.find("*** END OF")
    return text[text.find("\n", start) + 1:end].strip() if start >= 0 and end > start else text


def book_pdf(text: str, target: Path) -> None:
    pdf = FPDF()
    pdf.set_auto_page_break(True, 20)
    pdf.add_page()
    pdf.set_font("helvetica", size=11)
    plain = text.translate({0x2018: "'", 0x2019: "'", 0x201C: '"', 0x201D: '"', 0x2014: "-", 0x2013: "-"})
    pdf.multi_cell(0, 6, plain.encode("latin-1", "replace").decode("latin-1"))
    pdf.output(str(target))


def font(size: int) -> ImageFont.FreeTypeFont:
    for name in ("segoeui.ttf", "arial.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default(size)


def download_photos(folder: Path) -> None:
    client = httpx.Client(follow_redirects=True, timeout=60,
                          headers={"User-Agent": "Dowse-testbed/0.1 (https://github.com/; test data)"})
    for name, source in PHOTOS:
        target = folder / name
        if target.exists():
            continue
        response = client.get(f"https://commons.wikimedia.org/wiki/Special:FilePath/{source}", params={"width": 1024})
        if response.status_code != 200:
            print(f"  couldn't fetch {source}: {response.status_code}")
            continue
        image = Image.open(io.BytesIO(response.content)).convert("RGB")
        image.save(target, quality=90)
        time.sleep(0.5)


def screenshot(lines: list[str], target: Path) -> None:
    image = Image.new("RGB", (1600, 900), (245, 246, 248))
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1600, 60), fill=(32, 33, 36))
    draw.text((24, 14), lines[0], fill="white", font=font(26))
    y = 140
    for index, line in enumerate(lines[1:]):
        draw.text((120, y), line, fill=(20, 20, 20) if index else (200, 30, 30), font=font(44 if index == 0 else 30))
        y += 90 if index == 0 else 60
    image.save(target)


def text_pdf(sections: list[tuple[str, str]], target: Path) -> None:
    pdf = FPDF()
    pdf.set_auto_page_break(True, 20)
    for title, body in sections:
        pdf.add_page()
        pdf.set_font("helvetica", "B", 18)
        pdf.multi_cell(0, 10, title, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("helvetica", size=12)
        pdf.multi_cell(0, 7, body, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(target))


def scanned_pdf(target: Path, scratch: Path) -> None:
    """A page that is only a picture of text (no text layer): only OCR can read its words."""
    page = Image.new("L", (1240, 1754), 250)
    draw = ImageDraw.Draw(page)
    y = 200
    for line in SCAN_TEXT:
        draw.text((140, y), line, fill=30, font=font(48 if y == 200 else 36))
        y += 80
    picture = scratch / "scan.png"
    page.rotate(0.8, fillcolor=250).save(picture)
    pdf = FPDF()
    pdf.add_page()
    pdf.image(str(picture), x=0, y=0, w=210, h=297)
    pdf.output(str(target))
    picture.unlink()


def word_doc(target: Path) -> None:
    document = Document()
    document.core_properties.title = "Q3 budget meeting"
    document.add_heading("Minutes: Q3 budget meeting", 1)
    document.add_paragraph("Attendees: finance team and project leads. The marketing budget is cut by 15% to fund "
                           "two new engineering hires. Travel is limited to customer visits until December.")
    document.add_paragraph("Action: Sara sends the revised spreadsheet by Friday.")
    document.save(target)


def plain_docx(target: Path, paragraphs: list[str]) -> None:
    """A Word document with no title of its own: what it's about is only in its name."""
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    document.save(target)


def main() -> None:
    root = Path(sys.argv[1]) if len(sys.argv) > 1 else paths.home() / "Testbed"
    print(f"Testbed: {root}")
    for sub in ("photos", "screenshots", "pdf", "notes", "mail", "code/webapp", ".ssh", "secrets"):
        (root / sub).mkdir(parents=True, exist_ok=True)
    download_photos(root / "photos")
    screenshot(["Terminal", "webapp: npm start", "Error: listen EADDRINUSE: address already in use :::3000"],
               root / "code/webapp/Screenshot 2026-09-02 090000.png")
    for name, lines in SCREENSHOTS.items():
        screenshot(lines, root / "screenshots" / name)
    text_pdf(LEASE, root / "pdf/lease.pdf")
    text_pdf([INVOICE], root / "pdf/scan_0042.pdf")
    scanned_pdf(root / "pdf/doc_scan_001.pdf", root)
    word_doc(root / "notes/meeting.docx")
    for name, text in DOCS.items():
        (root / name).write_text(text, encoding="utf-8")
    (root / "audio").mkdir(exist_ok=True)
    (root / "video").mkdir(exist_ok=True)
    with tempfile.TemporaryDirectory() as scratch:
        scratch = Path(scratch)
        if not (root / "audio/episode 12.mp3").exists():
            moments = make_podcast(root / "audio/episode 12.mp3", scratch)
            print("  podcast topics at", moments)
        if not (root / "audio/New Recording 3.m4a").exists():
            memo, rate = speak(MEMO, scratch)
            encode_audio(memo, rate, root / "audio/New Recording 3.m4a", "aac")
        encode_audio(np.zeros(10 * 16000, np.int16), 16000, root / "audio/silence.wav", "pcm_s16le")
        if not (root / "video/VID_4001.mp4").exists():
            make_slideshow([root / "photos/IMG_1001.jpg", root / "photos/IMG_1003.jpg", root / "photos/IMG_1004.jpg"],
                           root / "video/VID_4001.mp4", scratch)
    (root / "books").mkdir(exist_ok=True)
    for name, number in BOOKS.items():
        if not (root / name).exists():
            text = gutenberg(number)
            if name.endswith(".pdf"):
                book_pdf(text, root / name)
            else:
                (root / name).write_text(text, encoding="utf-8")
    hard_testbed.make(root, text_pdf, plain_docx)
    print(f"{sum(1 for p in root.rglob('*') if p.is_file())} files")


if __name__ == "__main__":
    main()
